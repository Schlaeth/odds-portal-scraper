import types

import pytest
from openpyxl import load_workbook

from odds_portal_scraper.basketball import (
    export_basketball_schedule,
    export_basketball_season,
    export_basketball_season_range,
)


@pytest.fixture
def sample_html():
    return """
    <html>
      <body>
        <table id="totals">
          <thead>
            <tr><th>Team</th><th>Wins</th></tr>
          </thead>
          <tbody>
            <tr><td>Team A</td><td>50</td></tr>
          </tbody>
        </table>
        <table id="standings">
          <thead>
            <tr><th>Seed</th><th>Team</th></tr>
          </thead>
          <tbody>
            <tr><td>1</td><td>Team A</td></tr>
          </tbody>
        </table>
      </body>
    </html>
    """


def test_export_basketball_season_writes_xlsx(tmp_path, monkeypatch, sample_html):
    def fake_get(_url, timeout, **_):
        resp = types.SimpleNamespace(status_code=200, text=sample_html)

        def raise_for_status():
            return None

        resp.raise_for_status = raise_for_status
        return resp

    monkeypatch.setattr("requests.get", fake_get)

    output_path = tmp_path / "season.xlsx"
    export_basketball_season("NBA", 2024, output_path)

    assert output_path.exists()
    wb = load_workbook(output_path)
    assert set(wb.sheetnames) == {"totals", "standings"}


def test_export_basketball_range(tmp_path, monkeypatch, sample_html):
    calls = []

    def fake_get(url, timeout, **_):
        calls.append(url)
        resp = types.SimpleNamespace(status_code=200, text=sample_html)

        def raise_for_status():
            return None

        resp.raise_for_status = raise_for_status
        return resp

    monkeypatch.setattr("requests.get", fake_get)

    out_dir = tmp_path / "out"
    written = export_basketball_season_range("NBA", 2023, 2024, out_dir)

    assert len(written) == 2
    assert all(p.exists() for p in written)
    assert "NBA_2023" in calls[0]
    assert "NBA_2024" in calls[1]


def test_export_basketball_schedule_workbook(tmp_path, monkeypatch):
    schedule_table = """
    <html>
      <body>
        <table id="schedule">
          <thead><tr><th>Date</th><th>Visitor</th></tr></thead>
          <tbody>
            <tr><td>2024-10-01</td><td>Team A</td></tr>
            <tr><td>2024-11-05</td><td>Team B</td></tr>
          </tbody>
        </table>
      </body>
    </html>
    """

    season_table = """
    <html>
      <body>
        <table id="totals"><tbody><tr><td>Team</td></tr></tbody></table>
      </body>
    </html>
    """

    def fake_get(url, timeout, **_):
        resp = types.SimpleNamespace(status_code=200, text=schedule_table if "games" in url else season_table)

        def raise_for_status():
            return None

        resp.raise_for_status = raise_for_status
        return resp

    monkeypatch.setattr("requests.get", fake_get)

    out = tmp_path / "season.xlsx"
    schedule_out = tmp_path / "season-games.xlsx"
    export_basketball_season(
        "NBA", 2024, out, include_schedule=True, schedule_output_path=schedule_out, schedule_format="xlsx"
    )

    wb = load_workbook(schedule_out)
    assert set(wb.sheetnames) == {"October", "November"}


def test_export_basketball_schedule_csv(tmp_path, monkeypatch):
    schedule_table = """
    <html>
      <body>
        <table id="schedule">
          <thead><tr><th>Date</th><th>Visitor</th></tr></thead>
          <tbody><tr><td>2024-10-01</td><td>Team A</td></tr></tbody>
        </table>
      </body>
    </html>
    """

    def fake_get(url, timeout, **_):
        resp = types.SimpleNamespace(status_code=200, text=schedule_table)

        def raise_for_status():
            return None

        resp.raise_for_status = raise_for_status
        return resp

    monkeypatch.setattr("requests.get", fake_get)

    out = tmp_path / "sched.csv"
    export_basketball_schedule("NBA", 2024, out, schedule_format="csv")
    assert out.exists()
